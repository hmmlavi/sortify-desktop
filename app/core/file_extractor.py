"""Safe, bounded file content and metadata extraction with archive/image bomb protections."""

from __future__ import annotations

import csv
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from app.models.file_record import IMAGE_EXTENSIONS

MAX_FILE_BYTES = 100 * 1024 * 1024  # 100 MB
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024  # 100 MB zip-bomb cap
MAX_ARCHIVE_ENTRIES = 5000
MAX_IMAGE_PIXELS = 64_000_000  # 64 MP decompression-bomb cap


class ExtractionError(Exception):
    """Raised when a file cannot be safely read or parsed."""


class _SafeHTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"} and self._skip_depth > 0:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            cleaned = data.strip()
            if cleaned:
                self._chunks.append(cleaned)

    def get_text(self) -> str:
        return " ".join(self._chunks)


def _validate_zip_safety(path: Path) -> None:
    """Inspect ZIP-based container headers (.docx, .pptx, .xlsx) to prevent zip bombs."""
    try:
        with zipfile.ZipFile(path, "r") as zf:
            infos = zf.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise ExtractionError(
                    f"Archive contains too many entries ({len(infos)} > {MAX_ARCHIVE_ENTRIES})."
                )
            total_uncompressed = sum(max(0, info.file_size) for info in infos)
            if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ExtractionError(
                    "Archive uncompressed payload exceeds the 100 MB safety limit."
                )
    except zipfile.BadZipFile as e:
        raise ExtractionError("Corrupt or invalid Office document container.") from e


def _strip_rtf(raw: str) -> str:
    """Extract readable plain text from an RTF document string."""
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    text = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+", " ", text).strip()


def extract_content(path: Path, max_chars: int = 12000) -> tuple[str, str]:
    """Extract bounded text excerpt and content kind from a supported file."""
    path = Path(path)
    ext = path.suffix.lower()
    max_chars = max(256, min(100_000, int(max_chars or 12000)))

    try:
        if path.is_symlink():
            raise ExtractionError("Symbolic links are not processed for safety.")
        if not path.exists() or not path.is_file():
            raise ExtractionError("File does not exist or is not a regular file.")

        stat = path.stat()
        if stat.st_size > MAX_FILE_BYTES:
            raise ExtractionError("File exceeds the 100 MB safety limit for this version.")

        if ext in {".txt", ".md", ".rst", ".log", ".yaml", ".yml"}:
            with path.open("r", encoding="utf-8", errors="replace") as f:
                return f.read(max_chars), "text"

        if ext == ".rtf":
            with path.open("r", encoding="utf-8", errors="replace") as f:
                raw = f.read(max_chars * 3)
            return _strip_rtf(raw)[:max_chars], "document"

        if ext in {".html", ".htm"}:
            with path.open("r", encoding="utf-8", errors="replace") as f:
                raw = f.read(max_chars * 3)
            parser = _SafeHTMLTextExtractor()
            parser.feed(raw)
            return parser.get_text()[:max_chars], "document"

        if ext == ".json":
            with path.open("r", encoding="utf-8", errors="replace") as f:
                raw = f.read(max_chars * 2)
            try:
                parsed = json.loads(raw)
                if isinstance(parsed, dict):
                    keys = ", ".join(list(parsed.keys())[:25])
                    summary = f"JSON keys: {keys}\n{json.dumps(parsed, ensure_ascii=False)[:max_chars]}"
                    return summary[:max_chars], "data"
            except (ValueError, TypeError):
                return raw[:max_chars], "data"
            return raw[:max_chars], "data"

        if ext in {".csv", ".tsv"}:
            delimiter = "\t" if ext == ".tsv" else ","
            lines: list[str] = []
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
                for i, row in enumerate(csv.reader(f, delimiter=delimiter)):
                    cleaned_cells = [c.strip() for c in row[:30] if c and c.strip()]
                    if cleaned_cells:
                        lines.append(" | ".join(cleaned_cells))
                    if i >= 35:
                        break
            return "\n".join(lines)[:max_chars], "spreadsheet"

        if ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as e:
                raise ExtractionError(
                    "PDF support requires pypdf. Install requirements.txt."
                ) from e
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise ExtractionError("Password-protected PDF is not supported.")
            chunks: list[str] = []
            try:
                meta = reader.metadata
                if meta:
                    title = getattr(meta, "title", None) or meta.get("/Title")
                    subject = getattr(meta, "subject", None) or meta.get("/Subject")
                    if title and str(title).strip():
                        chunks.append(f"Title: {str(title).strip()}")
                    if subject and str(subject).strip():
                        chunks.append(f"Subject: {str(subject).strip()}")
            except (AttributeError, KeyError, TypeError, ValueError):
                chunks = []
            for page in reader.pages[:20]:
                chunks.append(page.extract_text() or "")
                if sum(map(len, chunks)) >= max_chars:
                    break
            text = "\n".join(chunks)[:max_chars]
            if not text.strip():
                return "", "pdf_scan_or_empty"
            return text, "pdf"

        if ext == ".docx":
            _validate_zip_safety(path)
            from docx import Document

            doc = Document(str(path))
            chunks = []
            try:
                if doc.core_properties.title:
                    chunks.append(f"Title: {doc.core_properties.title}")
                if doc.core_properties.subject:
                    chunks.append(f"Subject: {doc.core_properties.subject}")
            except (AttributeError, KeyError, TypeError, ValueError):
                chunks = []
            for p in doc.paragraphs:
                if p.text and p.text.strip():
                    chunks.append(p.text.strip())
                if sum(map(len, chunks)) >= max_chars:
                    break
            if sum(map(len, chunks)) < max_chars:
                for table in doc.tables[:5]:
                    for row in table.rows[:15]:
                        cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                        if cells:
                            chunks.append(" | ".join(cells))
            return "\n".join(chunks)[:max_chars], "document"

        if ext == ".pptx":
            _validate_zip_safety(path)
            from pptx import Presentation

            prs = Presentation(str(path))
            chunks = []
            # python-pptx Slides supports iteration but not slicing, so bound the
            # loop explicitly (slicing raises AttributeError on 'list').
            for index, slide in enumerate(prs.slides):
                if index >= 30:
                    break
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text and shape.text.strip():
                        chunks.append(shape.text.strip())
                if sum(map(len, chunks)) >= max_chars:
                    break
            return "\n".join(chunks)[:max_chars], "presentation"

        if ext == ".xlsx":
            _validate_zip_safety(path)
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True, data_only=True)
            chunks = []
            try:
                for sheet_index, ws in enumerate(wb.worksheets):
                    if sheet_index >= 10:
                        break
                    chunks.append(f"Sheet: {ws.title}")
                    for row in ws.iter_rows(max_row=30, max_col=20, values_only=True):
                        vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                        if vals:
                            chunks.append(" | ".join(vals))
                    if sum(map(len, chunks)) >= max_chars:
                        break
            finally:
                wb.close()
            return "\n".join(chunks)[:max_chars], "spreadsheet"

        if ext in IMAGE_EXTENSIONS:
            try:
                from PIL import Image

                Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
                with Image.open(path) as im:
                    if im.width * im.height > MAX_IMAGE_PIXELS:
                        raise ExtractionError(
                            "Image resolution exceeds the 64 MP safety limit."
                        )
                    aspect = "landscape" if im.width > im.height else ("portrait" if im.height > im.width else "square")
                    meta_parts = [
                        f"Image {im.format or ext.lstrip('.').upper()} {im.width}x{im.height} ({aspect}, {im.mode})"
                    ]
                    exif = im.getexif()
                    if exif:
                        # Extract safe descriptive EXIF tags (306=DateTime, 36867=DateTimeOriginal, 270=ImageDescription, 271=Make, 272=Model)
                        for tag_id, label in (
                            (270, "Description"),
                            (36867, "Captured"),
                            (306, "Date"),
                            (271, "Camera Make"),
                            (272, "Camera Model"),
                        ):
                            val = exif.get(tag_id)
                            if val and isinstance(val, str) and val.strip():
                                meta_parts.append(f"{label}: {val.strip()[:80]}")
                    return "; ".join(meta_parts), "image"
            except ImportError:
                return (
                    "Image file. Visual content requires a vision-capable AI provider; no image description available locally.",
                    "image",
                )

        raise ExtractionError(f'Unsupported extension: {ext or "(none)"}')
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Could not extract content: {type(e).__name__}: {e}") from e
