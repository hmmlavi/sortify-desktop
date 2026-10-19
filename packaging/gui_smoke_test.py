"""Headless GUI smoke test for the packaged (or source) Sortify application.

It exercises the real widgets (not a mock) against a temporary folder:

1. launch the main window and render it,
2. scan a disposable folder of mixed documents,
3. generate local filename suggestions,
4. edit a proposal and confirm collisions are detected and resolved,
5. apply the batch, verify files on disk and history rows,
6. undo the session and verify the original names are restored,
7. save screenshots of all three pages for visual review.

Run with::

    python packaging/gui_smoke_test.py --output build/gui-artifacts

Exit code is non-zero if any step fails.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _wait_for(predicate, timeout: float = 30.0, interval: float = 0.02, pump=None) -> bool:
    """Poll ``predicate`` while pumping the Qt event loop.

    Worker signals are queued to the GUI thread, so a plain ``sleep`` loop would
    never observe them.
    """
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if pump is not None:
            pump()
        if predicate():
            return True
        time.sleep(interval)
    if pump is not None:
        pump()
    return predicate()


class SmokeReport:
    def __init__(self) -> None:
        self.steps: list[dict] = []
        self.ok = True

    def check(self, name: str, condition: bool, detail: str = "") -> bool:
        self.steps.append({"step": name, "ok": bool(condition), "detail": detail})
        if not condition:
            self.ok = False
        print(
            f"[{'PASS' if condition else 'FAIL'}] {name}" + (f" - {detail}" if detail else ""),
            flush=True,
        )
        return bool(condition)


def create_sample_files(folder: Path) -> list[Path]:
    """Create small, real, disposable files of several supported types."""
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    from pptx import Presentation

    folder.mkdir(parents=True, exist_ok=True)
    created: list[Path] = []

    notes_a = folder / "notes copy.txt"
    notes_a.write_text(
        "Quarterly revenue review\n"
        "The Northwind account renewed its contract for another year. "
        "Pipeline coverage improved across the mid-market segment.\n",
        encoding="utf-8",
    )
    created.append(notes_a)

    notes_b = folder / "UNTITLED(1).txt"
    notes_b.write_text(
        "Quarterly revenue review\n"
        "The Northwind account renewed its contract for another year. "
        "Pipeline coverage improved across the mid-market segment.\n",
        encoding="utf-8",
    )
    created.append(notes_b)

    csv_like = folder / "data_export.csv"
    csv_like.write_text(
        "region,revenue,quarter\nEMEA,120000,Q3\nAPAC,98000,Q3\n", encoding="utf-8"
    )
    created.append(csv_like)

    doc = Document()
    doc.add_heading("Supplier onboarding policy", level=1)
    doc.add_paragraph(
        "This policy describes the standard onboarding checks for new suppliers, "
        "including compliance screening and payment terms."
    )
    docx_path = folder / "document final final.docx"
    doc.save(docx_path)
    created.append(docx_path)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Headcount"
    sheet.append(["team", "headcount", "cost"])
    sheet.append(["Design", 12, 840000])
    xlsx_path = folder / "Book1.xlsx"
    workbook.save(xlsx_path)
    created.append(xlsx_path)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Q4 launch readiness"
    slide.placeholders[1].text = "Timeline, owners, and open risks for the Q4 release."
    pptx_path = folder / "presentation1.pptx"
    presentation.save(pptx_path)
    created.append(pptx_path)

    image = Image.new("RGB", (320, 220), (32, 48, 44))
    image_path = folder / "IMG_0042.png"
    image.save(image_path)
    created.append(image_path)

    unsupported = folder / "archive.bin"
    unsupported.write_bytes(b"\x00\x01\x02not a document")
    created.append(unsupported)
    return created


def run(output_dir: Path) -> int:
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    data_dir = tempfile.mkdtemp(prefix="sortify-smoke-data-")
    os.environ["SORTIFY_DATA_DIR"] = data_dir

    from PySide6.QtWidgets import QApplication, QMessageBox

    # Modal dialogs cannot be answered in a headless run; record and auto-confirm
    # them so the real workflow code path still executes.
    message_log: list[tuple[str, str, str]] = []

    def _auto(kind: str, default):
        def handler(parent=None, title="", text="", *args, **kwargs):
            message_log.append((kind, str(title), str(text)))
            print(f"       dialog[{kind}] {title}: {text.splitlines()[0][:90]}", flush=True)
            return default()

        return staticmethod(handler)

    QMessageBox.information = _auto("information", lambda: QMessageBox.StandardButton.Ok)
    QMessageBox.warning = _auto("warning", lambda: QMessageBox.StandardButton.Yes)
    QMessageBox.critical = _auto("critical", lambda: QMessageBox.StandardButton.Ok)
    QMessageBox.question = _auto("question", lambda: QMessageBox.StandardButton.Yes)

    from app.config import Settings
    from app.core.operation_history import OperationHistory
    from app.ui.icons import get_brand_icon
    from app.ui.main_window import MainWindow

    report = SmokeReport()
    output_dir.mkdir(parents=True, exist_ok=True)
    app = QApplication.instance() or QApplication([sys.argv[0]])
    app.setStyle("Fusion")

    settings = Settings()
    settings.confirm_before_apply = False
    settings.cache_enabled = False
    history = OperationHistory(Path(data_dir) / "history.sqlite3")
    window = MainWindow(settings=settings, history=history)
    window.setWindowIcon(get_brand_icon())
    window.resize(1440, 900)
    window.show()
    app.processEvents()

    report.check("window visible", window.isVisible(), f"{window.width()}x{window.height()}")
    report.check(
        "brand icon available",
        not get_brand_icon().isNull(),
        f"sizes={[s.width() for s in get_brand_icon().availableSizes()]}",
    )
    report.check(
        "tagline free of retired branding",
        "spatial" not in window.windowTitle().lower(),
        window.windowTitle(),
    )

    shots: list[str] = []

    def grab(name: str) -> None:
        app.processEvents()
        path = output_dir / name
        window.grab().save(str(path))
        shots.append(str(path))
        print(f"       screenshot -> {path}", flush=True)

    grab("01-workspace-empty.png")

    workdir = Path(tempfile.mkdtemp(prefix="sortify-smoke-files-"))
    files = create_sample_files(workdir)
    original_bytes = {p: p.read_bytes() for p in files}
    report.check("sample files created", len(files) == 8, f"{len(files)} files")

    window.folder = workdir
    window._set_folder_labels()
    window.start_scan()
    scanned = _wait_for(
        lambda: window._active_worker is None and bool(window.records), pump=app.processEvents
    )
    app.processEvents()
    report.check("scan completed", scanned, f"{len(window.records)} records")
    report.check(
        "unsupported file detected",
        any(r.status == "Unsupported" for r in window.records),
    )
    grab("02-workspace-scanned.png")

    window.start_analysis()
    analyzed = _wait_for(
        lambda: window._active_worker is None
        and any(r.proposed_name for r in window.records),
        timeout=60,
        pump=app.processEvents,
    )
    app.processEvents()
    proposed = [r for r in window.records if r.proposed_name]
    report.check("names generated", analyzed and bool(proposed), f"{len(proposed)} proposals")
    report.check(
        "all proposals keep their extension",
        all(r.proposed_name.lower().endswith(r.extension.lower()) for r in proposed),
    )
    # Prefer a row that actually has a proposal for the visual captures.
    preferred_row = next(
        (index for index, record in enumerate(window.records) if record.proposed_name), 0
    )
    if window.table.rowCount():
        window.table.selectRow(preferred_row)
        app.processEvents()
    grab("03-review.png")

    # Inspector round-trip: render the selected record in the inspector.
    if window.table.rowCount():
        window.table.selectRow(preferred_row)
        app.processEvents()
        report.check(
            "inspector populated",
            window.inspector_filename.text() not in {"", "Select a file"},
            window.inspector_filename.text(),
        )
        grab("04-review-inspector.png")

    # Manual edit through the inspector editor.
    if window.records:
        target = window.records[0]
        window._selected_record_path = str(target.path)
        window.inspector_name_edit.setText("Edited By Smoke Test")
        window._inspector_name_committed()
        app.processEvents()
        report.check(
            "manual edit applied",
            target.proposed_name == "Edited_By_Smoke_Test" + target.extension,
            target.proposed_name,
        )

    # Force a duplicate to exercise conflict handling.
    editable = [r for r in window.records if r.proposed_name]
    if len(editable) >= 2:
        for record in editable:
            record.proposed_name = "Collision" + record.extension
        window._recompute_collisions()
        forced = sum(1 for r in editable if r.conflict_reason)
        report.check("duplicate conflicts detected", forced >= 2, f"{forced} conflicts")
        window.resolve_batch_conflicts()
        app.processEvents()
        resolved_names = [r.proposed_name for r in editable]
        report.check(
            "conflicts resolved uniquely",
            len({n.casefold() for n in resolved_names}) == len(resolved_names),
            ", ".join(resolved_names[:4]),
        )

    report.check(
        "file contents untouched before apply",
        all(p.read_bytes() == original_bytes[p] for p in files if p.exists()),
    )

    window._execute_rename_batch(window._collect_selected_proposals())
    app.processEvents()
    time.sleep(0.2)
    app.processEvents()

    history_rows = history.recent(100)
    succeeded = [row for row in history_rows if row["status"] == "success"]
    report.check("history rows recorded", len(succeeded) >= 1, f"{len(succeeded)} rows")
    report.check(
        "renamed files exist on disk",
        all(Path(row["new_path"]).exists() for row in succeeded),
        "",
    )
    report.check(
        "contents preserved after rename",
        all(
            Path(row["new_path"]).read_bytes() == original_bytes[Path(row["original_path"])]
            for row in succeeded
            if Path(row["original_path"]) in original_bytes
        ),
        "byte comparison",
    )

    window._navigate(2)
    app.processEvents()
    grab("05-history.png")

    undo = window.engine.undo_last_session()
    app.processEvents()
    restored = len(undo.get("success", []))
    report.check("undo restored files", restored >= 1, f"{restored} restored")
    report.check(
        "original names back on disk",
        all(Path(row["original_path"]).exists() for row in succeeded),
        "",
    )
    report.check(
        "contents still identical after undo",
        all(
            Path(row["original_path"]).read_bytes() == original_bytes[Path(row["original_path"])]
            for row in succeeded
            if Path(row["original_path"]) in original_bytes
        ),
        "byte comparison",
    )

    window.apply_theme(False)
    app.processEvents()
    grab("06-light-theme.png")
    window.apply_theme(True)
    app.processEvents()

    # Dialog rendering
    from app.ui.components import BatchPreviewDialog
    from app.ui.settings_dialog import SettingsDialog

    preview = BatchPreviewDialog(
        [(files[0], "Proposed_Name.txt"), (files[1], "Another_Name.txt")],
        dark_mode=True,
        parent=window,
    )
    preview.show()
    app.processEvents()
    preview.grab().save(str(output_dir / "07-preview-dialog.png"))
    shots.append(str(output_dir / "07-preview-dialog.png"))
    preview.close()

    settings_dialog = SettingsDialog(window.settings, window, history=history)
    settings_dialog.show()
    app.processEvents()
    settings_dialog.grab().save(str(output_dir / "08-settings-dialog.png"))
    shots.append(str(output_dir / "08-settings-dialog.png"))
    settings_dialog.close()

    report.check("sidecar files cleaned up", not Path(data_dir, "history.sqlite3-wal").exists())

    window.close()
    app.processEvents()
    history.close()

    payload = {
        "ok": report.ok,
        "steps": report.steps,
        "screenshots": shots,
        "dialogs": message_log,
    }
    (output_dir / "smoke-report.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(json.dumps({"ok": report.ok, "steps": len(report.steps)}, indent=2))
    return 0 if report.ok else 1


def main() -> int:
    # Windows consoles default to a legacy code page; keep the harness output
    # ASCII-only and force UTF-8 for anything else.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Sortify GUI smoke test")
    parser.add_argument("--output", default="build/gui-artifacts")
    args = parser.parse_args()
    return run(Path(args.output).resolve())


if __name__ == "__main__":
    sys.exit(main())
